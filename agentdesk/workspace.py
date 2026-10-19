import base64
import io
import re
import uuid
from pathlib import Path

from PIL import Image

MAX_FILE = 20 * 1024 * 1024
TEXT_LIMIT = 60_000
SKIP = {".git", ".venv", "node_modules", "__pycache__", ".agentdesk", ".env"}

def resolve_inside(root: str | Path, relative: str, *, write=False) -> Path:
    root = Path(root).resolve(strict=True)
    path = (root / relative).resolve()
    if not path.is_relative_to(root) or path == root and write:
        raise ValueError("路径必须位于当前工作区内。")
    parts = path.relative_to(root).parts
    if any(p in SKIP or p.startswith(".env") for p in parts):
        raise ValueError("此路径属于配置或内部目录，不能通过文件工具访问。")
    return path

def extract_text(path: Path, limit=TEXT_LIMIT) -> str:
    if path.stat().st_size > MAX_FILE:
        raise ValueError("文件超过 20 MB。")
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        chunks, count = [], 0
        for i, page in enumerate(reader.pages[:100]):
            text = f"\n[第 {i+1} 页]\n{page.extract_text() or ''}"
            chunks.append(text)
            count += len(text)
            if count >= limit:
                break
        result = "".join(chunks)
        if not any(c.strip() for c in re.split(r"\[第 \d+ 页\]", result)):
            return "此 PDF 没有可提取的文本，可能是扫描件；请上传页面图片并使用视觉模型。"
    elif ext == ".docx":
        from docx import Document
        doc = Document(path)
        result = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            result += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    else:
        raw = path.read_bytes()
        if b"\x00" in raw[:8192]:
            return "此文件是二进制文件，已保存，不能直接提取正文。"
        try:
            result = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                result = raw.decode("gb18030")
            except UnicodeDecodeError:
                return "无法识别此文件的文本编码，文件已保存。"
    return result[:limit] + ("\n[正文过长，已截断]" if len(result) > limit else "")

def image_data(path: Path) -> str:
    with Image.open(path) as img:
        img.thumbnail((2048, 2048))
        img = img.convert("RGB")
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=88)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()

def save_upload(root: str, name: str, data: bytes) -> dict:
    if len(data) > MAX_FILE:
        raise ValueError(f"{name} 超过 20 MB。")
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)[-160:] or "file"
    relative = f"uploads/{uuid.uuid4().hex[:10]}_{name}"
    path = resolve_inside(root, relative, write=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    is_image = False
    if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}:
        try:
            with Image.open(path) as img:
                img.verify()
            is_image = True
        except Exception:
            pass
    return {"name": name, "path": relative, "size": len(data), "image": is_image}

def message_content(text: str, attachments: list, workspace: str, vision: bool):
    content = [{"type": "text", "text": text or "请分析我上传的附件。"}]
    budget = TEXT_LIMIT
    for item in attachments:
        path = resolve_inside(workspace, item["path"])
        label = f"\n附件：{item['name']}，工作区相对路径：{item['path']}\n"
        if item["image"] and vision:
            content.append({"type": "text", "text": label})
            content.append({"type": "image", "base64": image_data(path).split(",", 1)[1], "mime_type": "image/jpeg"})
        elif item["image"]:
            content.append({"type": "text", "text": label + "视觉输入未启用，不能声称已经看到了图像内容。"})
        else:
            try:
                extracted = extract_text(path, min(budget, 30_000)) if budget else "[附件正文总量超限，请用 read_file 阅读]"
            except Exception as exc:
                extracted = f"提取失败：{type(exc).__name__}；可以尝试工具读取。"
            budget = max(0, budget - len(extracted))
            content.append({"type": "text", "text": label + "<attachment_data>\n" + extracted + "\n</attachment_data>"})
    return content

def list_directory(root, relative="."):
    path = resolve_inside(root, relative)
    items = []
    for child in sorted(path.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower())):
        try:
            safe = resolve_inside(root, str(child.relative_to(Path(root).resolve())))
            items.append({"name": child.name, "path": child.relative_to(Path(root).resolve()).as_posix(),
                          "directory": safe.is_dir(), "size": safe.stat().st_size if safe.is_file() else 0})
        except (ValueError, OSError):
            continue
        if len(items) >= 300:
            break
    return items
